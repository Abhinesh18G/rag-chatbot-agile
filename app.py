import os
import json
import re
import uuid
import html as html_lib
from datetime import datetime
from io import BytesIO

import streamlit as st
from dotenv import load_dotenv
from pptx import Presentation

from langchain_google_genai import (
    GoogleGenerativeAIEmbeddings,
    ChatGoogleGenerativeAI,
)
from langchain_community.vectorstores import Chroma

from utils import extract_text_from_pdf, chunk_text, build_prompt

# --------------------------------------------------
# Page setup
# --------------------------------------------------
st.set_page_config(
    page_title="Intelligent RAG Study Hub",
    page_icon="📚",
    layout="wide",
)


def H(markup: str) -> str:
    """Flatten HTML so Markdown never treats indented lines as code blocks."""
    return " ".join(line.strip() for line in markup.splitlines() if line.strip())


def render(markup: str):
    st.markdown(H(markup), unsafe_allow_html=True)


# --------------------------------------------------
# Aurora Glass styling
# --------------------------------------------------
st.markdown(
    """
    <style>
    :root {
        --cyan: #69f0df;
        --blue: #5b9bff;
        --violet: #b57cff;
        --glass-border: rgba(150,190,255,.20);
        --panel: rgba(12,24,52,.62);
    }

    .stApp {
        background:
          radial-gradient(ellipse at 10% 5%, rgba(39,208,190,.12), transparent 35%),
          radial-gradient(ellipse at 90% 12%, rgba(124,92,255,.20), transparent 34%),
          radial-gradient(ellipse at 50% 105%, rgba(40,126,255,.16), transparent 40%),
          linear-gradient(135deg, #050c1c 0%, #0a1024 48%, #0f0d28 100%);
        background-attachment: fixed;
    }

    header[data-testid="stHeader"] { background: transparent; }

    .block-container {
        padding-top: 1.4rem;
        padding-bottom: 3rem;
        max-width: 1280px;
    }

    h1, h2, h3 { letter-spacing: -.02em; }

    /* ---------- Sidebar ---------- */
    [data-testid="stSidebar"] {
        background: linear-gradient(180deg, rgba(7,16,38,.96), rgba(12,12,40,.96));
        border-right: 1px solid rgba(120,150,255,.25);
        box-shadow: 4px 0 40px rgba(90,70,255,.12);
    }
    [data-testid="stSidebar"] .block-container { padding-top: 1rem; }

    .brand { text-align: center; padding: .6rem 0 1rem; }
    .brand-logo {
        width: 54px; height: 54px; margin: 0 auto .6rem;
        border-radius: 16px; display: flex; align-items: center;
        justify-content: center; font-size: 1.8rem;
        background: linear-gradient(135deg, rgba(105,240,223,.25), rgba(124,92,255,.3));
        border: 1px solid rgba(150,210,255,.4);
        box-shadow: 0 0 24px rgba(105,240,223,.25);
    }
    .brand-name { font-weight: 800; font-size: 1.12rem; line-height: 1.25; color: #fff; }

    .side-label {
        font-size: .68rem; letter-spacing: .18em; font-weight: 700;
        color: #8fa2d6; margin: 1.1rem 0 .3rem .4rem;
    }

    [data-testid="stSidebar"] .stButton > button {
        justify-content: flex-start;
        text-align: left;
        min-height: 2.8rem;
        border-radius: 14px;
        background: transparent;
        border: 1px solid transparent;
        box-shadow: none;
        color: #d4e0ff;
        font-weight: 550;
        padding-left: .9rem;
    }
    [data-testid="stSidebar"] .stButton > button:hover {
        background: rgba(120,160,255,.10);
        border-color: rgba(120,170,255,.25);
        transform: none;
        box-shadow: none;
    }
    [data-testid="stSidebar"] button[data-testid="stBaseButton-primary"],
    [data-testid="stSidebar"] button[kind="primary"] {
        background: linear-gradient(90deg, rgba(48,120,255,.60), rgba(60,200,230,.20));
        border: 1px solid rgba(130,200,255,.75);
        color: #fff;
        box-shadow: 0 0 22px rgba(70,150,255,.35), inset 0 1px 0 rgba(255,255,255,.15);
    }

    .side-card {
        margin-top: 1.5rem; padding: 1rem 1.1rem; border-radius: 16px;
        border: 1px solid rgba(160,120,255,.35);
        background: linear-gradient(135deg, rgba(60,200,150,.10), rgba(130,90,255,.14));
        font-size: .92rem; line-height: 1.5; color: #dfe7ff;
    }
    .side-card b { color: #fff; }

    /* ---------- Top bar ---------- */
    .topbar { display: flex; justify-content: flex-end; margin-bottom: .5rem; }
    .welcome {
        display: flex; align-items: center; gap: .65rem;
        padding: .45rem .9rem .45rem .5rem; border-radius: 99px;
        border: 1px solid var(--glass-border); background: var(--panel);
        font-size: .82rem; line-height: 1.2;
    }
    .welcome .avatar {
        width: 34px; height: 34px; border-radius: 50%;
        display: flex; align-items: center; justify-content: center;
        background: linear-gradient(135deg, #69f0df, #8a6bff); font-size: 1.05rem;
    }
    .welcome small { display: block; opacity: .65; font-size: .72rem; }

    /* ---------- Hero ---------- */
    .hero {
        position: relative; overflow: hidden;
        padding: 2.2rem 2.4rem; border-radius: 26px; margin-bottom: 1.3rem;
        background:
          radial-gradient(ellipse at 75% 40%, rgba(70,110,255,.28), transparent 55%),
          linear-gradient(115deg, rgba(20,60,110,.35), rgba(30,25,90,.45));
        border: 1px solid rgba(140,190,255,.28);
        box-shadow: 0 20px 70px rgba(0,0,0,.35), inset 0 1px 0 rgba(255,255,255,.07);
        animation: rise-in .65s ease-out both;
    }
    .hero-grid {
        display: grid; grid-template-columns: 1.15fr .85fr;
        gap: 1rem; align-items: center;
    }
    .hero-kicker {
        display: inline-block; font-size: .72rem; letter-spacing: .22em;
        color: #b9c8f5; font-weight: 700; padding: .55rem 1rem;
        border-radius: 12px; border: 1px solid rgba(150,190,255,.3);
        background: rgba(10,20,50,.5); margin-bottom: 1.1rem;
    }
    .hero h1 {
        margin: 0 0 .9rem; padding: 0; line-height: 1.04; font-weight: 800;
        font-size: clamp(2.2rem, 5vw, 3.7rem);
    }
    .grad  { background: linear-gradient(90deg, #5ff0e0, #5aa4ff);
             -webkit-background-clip: text; background-clip: text; color: transparent; }
    .grad2 { background: linear-gradient(90deg, #5aa4ff, #c58cff);
             -webkit-background-clip: text; background-clip: text; color: transparent; }
    .hero p { margin: 0 0 1.3rem; opacity: .88; font-size: 1.05rem;
              max-width: 520px; line-height: 1.6; }

    .chips { display: flex; flex-wrap: wrap; gap: .6rem; }
    .chip {
        display: inline-flex; align-items: center; gap: .5rem;
        padding: .5rem .85rem; border-radius: 12px; font-size: .82rem;
        border: 1px solid rgba(150,190,255,.28); background: rgba(10,22,52,.6);
    }

    .ai-card {
        position: absolute; right: 1.6rem; top: 1.6rem; width: 205px;
        padding: .9rem 1.05rem; border-radius: 16px; font-size: .82rem;
        line-height: 1.45; border: 1px solid rgba(105,240,223,.4);
        background: rgba(8,28,48,.72); backdrop-filter: blur(12px);
        box-shadow: 0 0 30px rgba(60,200,220,.15);
    }
    .ai-card .ai-title { color: var(--cyan); font-weight: 800;
        letter-spacing: .1em; display: flex; justify-content: space-between;
        align-items: center; margin-bottom: .35rem; }
    .dot { width: 11px; height: 11px; border-radius: 50%; background: #2fe38b;
           box-shadow: 0 0 12px #2fe38b; }

    /* Orb */
    .orb-wrap { position: relative; width: 320px; height: 320px; margin: 0 auto; }
    .orb {
        position: absolute; inset: 38px 38px 70px 38px; border-radius: 50%;
        background:
          radial-gradient(circle at 35% 28%, rgba(255,255,255,.95), rgba(120,245,235,.65) 14%,
                          rgba(60,130,255,.6) 42%, rgba(110,50,220,.55) 72%, rgba(18,10,60,.85));
        box-shadow: 0 0 60px rgba(70,170,255,.6), 0 0 130px rgba(150,90,255,.4),
                    inset 0 0 40px rgba(255,255,255,.25);
        animation: float 6s ease-in-out infinite alternate;
        overflow: hidden;
    }
    .orb:before {
        content: ""; position: absolute; inset: 0; border-radius: 50%;
        background:
          repeating-linear-gradient(35deg,  transparent 0 22px, rgba(255,255,255,.22) 22px 23px),
          repeating-linear-gradient(-35deg, transparent 0 22px, rgba(255,255,255,.18) 22px 23px);
        mix-blend-mode: screen;
    }
    .orb-ring {
        position: absolute; left: 6px; right: 6px; bottom: 14px; height: 70px;
        border-radius: 50%; border: 2px solid rgba(105,240,223,.7);
        box-shadow: 0 0 24px rgba(105,240,223,.55), inset 0 0 24px rgba(181,124,255,.5);
        background: radial-gradient(ellipse, rgba(181,124,255,.35), transparent 70%);
    }
    .orb-ring.two { left: 36px; right: 36px; bottom: 26px; height: 46px;
                    border-color: rgba(181,124,255,.8); }
    .spark { position: absolute; width: 9px; height: 9px; border-radius: 50%;
             background: #ff7be5; box-shadow: 0 0 14px #ff7be5;
             animation: twinkle 3s ease-in-out infinite alternate; }
    .spark.a { left: 6%;  top: 62%; }
    .spark.b { right: 8%; top: 30%; background: #6fe9ff; box-shadow: 0 0 14px #6fe9ff;
               animation-delay: 1s; }
    .spark.c { right: 14%; bottom: 20%; animation-delay: 2s; }

    /* ---------- Panels ---------- */
    .st-key-upload_card, .st-key-docs_card {
        border: 1px solid rgba(110,170,255,.35); border-radius: 20px;
        background: linear-gradient(135deg, rgba(14,30,66,.75), rgba(14,22,56,.7));
        padding: 1.2rem 1.35rem;
        box-shadow: 0 14px 40px rgba(0,0,0,.25), inset 0 1px 0 rgba(255,255,255,.05);
    }
    .panel-head { display: flex; align-items: center; gap: .9rem; margin-bottom: .8rem; }
    .panel-ico {
        width: 48px; height: 48px; border-radius: 14px; display: flex;
        align-items: center; justify-content: center; font-size: 1.5rem;
        background: linear-gradient(135deg, rgba(80,140,255,.4), rgba(150,90,255,.4));
        border: 1px solid rgba(150,200,255,.35);
    }
    .panel-title { font-weight: 750; font-size: 1.12rem; }
    .panel-sub { opacity: .75; font-size: .88rem; }
    .panel-row { display: flex; justify-content: space-between; align-items: center; }

    .doc-row {
        display: flex; align-items: center; gap: .85rem;
        padding: .65rem 0; border-bottom: 1px solid rgba(150,190,255,.10);
    }
    .doc-row:last-child { border-bottom: 0; }
    .doc-ico {
        width: 44px; height: 48px; border-radius: 10px; display: flex;
        align-items: center; justify-content: center; font-weight: 800;
        font-size: .78rem; color: #fff; flex-shrink: 0;
    }
    .doc-ico.pdf  { background: linear-gradient(135deg, #ff5d5d, #c62b4a); }
    .doc-ico.pptx { background: linear-gradient(135deg, #ff9a4d, #e2552b); }
    .doc-meta { flex: 1; min-width: 0; }
    .doc-meta b { display: block; font-size: .93rem; overflow: hidden;
                  text-overflow: ellipsis; white-space: nowrap; }
    .doc-meta span { font-size: .78rem; opacity: .65; }
    .pill-ok {
        font-size: .78rem; font-weight: 600; padding: .3rem .7rem; border-radius: 99px;
        color: #54f0a6; background: rgba(40,200,120,.14);
        border: 1px solid rgba(60,220,140,.35); white-space: nowrap;
    }
    .empty-docs { opacity: .7; padding: 1.3rem .2rem; font-size: .92rem; line-height: 1.6; }

    .supported { display: flex; flex-wrap: wrap; gap: .5rem; align-items: center;
                 font-size: .85rem; margin-top: .6rem; opacity: .9; }
    .tag { padding: .2rem .65rem; border-radius: 8px; font-weight: 650; font-size: .78rem;
           background: rgba(80,120,255,.2); border: 1px solid rgba(120,160,255,.4); }

    /* Feature cards */
    .st-key-card_quiz, .st-key-card_exam {
        border-radius: 22px; padding: 1.25rem 1.4rem;
        box-shadow: 0 16px 44px rgba(0,0,0,.28);
    }
    .st-key-card_quiz {
        border: 1px solid rgba(80,170,255,.5);
        background: linear-gradient(135deg, rgba(14,50,100,.65), rgba(12,30,72,.7));
    }
    .st-key-card_exam {
        border: 1px solid rgba(170,110,255,.5);
        background: linear-gradient(135deg, rgba(60,28,110,.62), rgba(30,22,76,.7));
    }
    .feat { display: flex; gap: 1.1rem; align-items: center; margin-bottom: .9rem; }
    .feat-ico {
        width: 74px; height: 74px; border-radius: 50%; flex-shrink: 0;
        display: flex; align-items: center; justify-content: center; font-size: 2.1rem;
        border: 2px solid rgba(110,200,255,.7);
        background: radial-gradient(circle, rgba(60,160,255,.35), rgba(10,30,70,.6));
        box-shadow: 0 0 28px rgba(70,170,255,.45);
    }
    .feat-ico.violet { border-color: rgba(190,130,255,.8);
        background: radial-gradient(circle, rgba(160,90,255,.4), rgba(30,15,70,.6));
        box-shadow: 0 0 28px rgba(160,100,255,.5); }
    .feat h3 { margin: 0 0 .25rem; font-size: 1.35rem; }
    .feat p { margin: 0; opacity: .82; font-size: .95rem; line-height: 1.5; }
    .feat-tags { display: flex; flex-wrap: wrap; gap: .5rem; margin-bottom: .6rem; }
    .feat-tag { padding: .4rem .75rem; border-radius: 12px; font-size: .78rem;
                border: 1px solid rgba(150,190,255,.25); background: rgba(8,16,40,.55); }

    /* ---------- Generic widgets ---------- */
    .hint-card {
        padding: 1rem 1.15rem; border-radius: 16px; margin: .55rem 0 1rem;
        border: 1px solid var(--glass-border);
        background: linear-gradient(120deg, rgba(105,240,223,.065),
                    rgba(124,167,255,.07), rgba(189,140,255,.055));
    }

    div.stButton > button,
    div.stDownloadButton > button,
    div.stFormSubmitButton > button {
        border-radius: 14px; min-height: 3rem; font-weight: 650;
        border: 1px solid rgba(145,188,255,.27);
        transition: transform .18s ease, box-shadow .22s ease,
                    border-color .22s ease, filter .22s ease;
        box-shadow: 0 5px 18px rgba(0,0,0,.12);
    }
    div.stButton > button:hover,
    div.stDownloadButton > button:hover,
    div.stFormSubmitButton > button:hover {
        transform: translateY(-2px); border-color: rgba(105,240,223,.65);
        box-shadow: 0 8px 28px rgba(44,207,202,.16); filter: brightness(1.08);
    }
    div.stButton > button:focus-visible,

    div.stDownloadButton > button:focus-visible {
        outline: 2px solid var(--cyan);
        outline-offset: 3px;
    }

    div.stButton > button[data-testid="stBaseButton-primary"],
    div.stFormSubmitButton > button[data-testid="stBaseButton-primaryFormSubmit"] {
        background: linear-gradient(90deg, #2f7bff, #8a5bff);
        border-color: rgba(160,200,255,.5);
    }

    div[data-testid="stFileUploader"] section {
        border: 1.5px dashed rgba(105,240,223,.6); border-radius: 16px;
        background: rgba(8,20,48,.5); transition: border-color .2s, background .2s;
    }
    div[data-testid="stFileUploader"] section:hover {
        border-color: rgba(189,140,255,.85); background: rgba(124,167,255,.09);
    }
    div[data-testid="stFileUploader"] button {
        border-radius: 12px; border: 1px solid rgba(110,180,255,.6);
        background: rgba(40,110,255,.25);
    }

    div[data-testid="stVerticalBlockBorderWrapper"] {
        border-color: var(--glass-border) !important; border-radius: 19px !important;
        background: rgba(255,255,255,.025); box-shadow: 0 12px 32px rgba(0,0,0,.10);
    }
    div[data-testid="stMetric"] {
        padding: .9rem 1rem; border: 1px solid var(--glass-border); border-radius: 16px;
        background: linear-gradient(135deg, rgba(105,240,223,.055), rgba(124,167,255,.055));
    }
    div[data-testid="stProgressBar"] > div { border-radius: 99px; overflow: hidden; }
    div[data-testid="stProgressBar"] > div > div {
        background: linear-gradient(90deg, #4ee4cf, #729bff, #c18aff);
        background-size: 200% 100%; animation: progress-shimmer 3s linear infinite;
    }
    div[data-testid="stExpander"] {
        border: 1px solid rgba(181,203,255,.16); border-radius: 15px;
        background: rgba(255,255,255,.018);
    }
    div[data-testid="stRadio"] label { transition: color .15s ease; }
    div[data-testid="stRadio"] label:hover { color: var(--cyan); }
    hr { border-color: rgba(181,203,255,.14); }

    @keyframes rise-in { from { opacity: 0; transform: translateY(9px); }
                         to { opacity: 1; transform: translateY(0); } }
    @keyframes float { from { transform: translateY(-6px); } to { transform: translateY(8px); } }
    @keyframes twinkle { from { opacity: .35; transform: scale(.8); }
                         to { opacity: 1; transform: scale(1.25); } }
    @keyframes progress-shimmer { 0% { background-position: 0% 50%; }
                                  100% { background-position: 200% 50%; } }

    @media (prefers-reduced-motion: reduce) {
        *,
        *:before,
        *:after {
            animation-duration: .01ms !important;
            animation-iteration-count: 1 !important;
            transition-duration: .01ms !important;
        }
    }
    @media (max-width: 1000px) {
        .hero-grid { grid-template-columns: 1fr; }
        .orb-wrap { display: none; }
        .ai-card { position: static; width: auto; margin-top: 1.2rem; }
    }
    @media (max-width: 700px) {
        .block-container { padding-left: 1rem; padding-right: 1rem; }
        .hero { padding: 1.4rem; border-radius: 18px; }
    }
    </style>
    """,
    unsafe_allow_html=True,
)

# --------------------------------------------------
# API and models
# --------------------------------------------------
load_dotenv()
GOOGLE_API_KEY = os.getenv("GOOGLE_API_KEY")

if not GOOGLE_API_KEY:
    st.error("GOOGLE_API_KEY was not found. Add it to your local .env file.")
    st.stop()


@st.cache_resource
def get_embedding_model():
    return GoogleGenerativeAIEmbeddings(
        model="models/gemini-embedding-001",
        google_api_key=GOOGLE_API_KEY,
    )


@st.cache_resource
def get_llm():
    return ChatGoogleGenerativeAI(
        model="gemini-3.6-flash",
        temperature=0,
        google_api_key=GOOGLE_API_KEY,
    )


embedding_model = get_embedding_model()
llm = get_llm()

# --------------------------------------------------
# Session state
# --------------------------------------------------
defaults = {
    "page": "Home",
    "retriever": None,
    "vector_db": None,
    "processed_files": None,
    "documents": [],
    "top_k": 4,
    "generated_quiz": None,
    "quiz_submitted": False,
    "quiz_answers": {},
    "quiz_topic_used": "",
    "quiz_history": [],
    "last_exam_answer": None,
    "activity": [],
}

for key, value in defaults.items():
    if key not in st.session_state:
        st.session_state[key] = value


def log_activity(message: str):
    st.session_state.activity.insert(
        0, {"time": datetime.now().strftime("%d %b, %I:%M %p"), "text": message}
    )
    del st.session_state.activity[50:]


def go(page: str):
    st.session_state.page = page


def human_size(num_bytes: int) -> str:
    size = float(num_bytes)
    for unit in ("B", "KB", "MB", "GB"):
        if size < 1024 or unit == "GB":
            return f"{size:.0f} {unit}" if unit == "B" else f"{size:.1f} {unit}"
        size /= 1024


# --------------------------------------------------
# PowerPoint extraction
# --------------------------------------------------
def extract_text_from_pptx(file_bytes):
    presentation = Presentation(BytesIO(file_bytes))
    slide_text = []

    for slide_number, slide in enumerate(presentation.slides, start=1):
        parts = [f"Slide {slide_number}"]

        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text.strip():
                parts.append(shape.text.strip())

            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    parts.append(" | ".join(cell.text.strip() for cell in row.cells))

        slide_text.append("\n".join(parts))

    return "\n\n".join(slide_text)


# --------------------------------------------------
# Sidebar navigation
# --------------------------------------------------
NAV_MAIN = [
    ("Home", ":material/home:"),
    ("Upload Documents", ":material/upload_file:"),
    ("Ask from Notes", ":material/chat:"),
    ("Interactive Quiz", ":material/sports_esports:"),
    ("10-Mark Preparation", ":material/description:"),
]
NAV_TOOLS = [
    ("My Study Progress", ":material/bar_chart:"),
    ("Recent Activity", ":material/schedule:"),
    ("Settings", ":material/settings:"),
]


def nav_button(label: str, icon: str):
    st.button(
        label,
        key=f"nav_{label}",
        use_container_width=True,
        type="primary" if st.session_state.page == label else "secondary",
        on_click=go,
        args=(label,),
    )


with st.sidebar:
    render(
        """
        <div class="brand">
          <div class="brand-logo">📖</div>
          <div class="brand-name">Intelligent RAG<br/>Study Hub</div>
        </div>
        """
    )
    for label, icon in NAV_MAIN:
        nav_button(label, icon)

    render('<div class="side-label">STUDY TOOLS</div>')
    for label, icon in NAV_TOOLS:
        nav_button(label, icon)

    render(
        """
        <div class="side-card">🌱 <b>Learn today</b><br/>
        Build a smarter you tomorrow</div>
        """
    )

# --------------------------------------------------
# Document processing
# --------------------------------------------------


def build_retriever(vector_db):
    return vector_db.as_retriever(search_kwargs={"k": st.session_state.top_k})


def process_files(uploaded_files):
    signature = tuple((f.name, f.size) for f in uploaded_files)

    if signature == st.session_state.processed_files:
        return

    all_chunks = []
    doc_info = []
    progress = st.progress(0, text="Preparing your documents...")

    for i, uploaded_file in enumerate(uploaded_files):
        progress.progress(
            int((i / max(len(uploaded_files), 1)) * 70),
            text=f"Reading {uploaded_file.name}...",
        )

        try:
            file_bytes = uploaded_file.getvalue()

            if uploaded_file.name.lower().endswith(".pdf"):
                text = extract_text_from_pdf(file_bytes)
            else:
                text = extract_text_from_pptx(file_bytes)

        except Exception:
            progress.empty()
            st.error(
                f"Could not read **{uploaded_file.name}**. "
                "Please check that the file is valid."
            )
            st.stop()

        if text.strip():
            chunks = chunk_text(text)
            all_chunks.extend(chunks)
            doc_info.append(
                {
                    "name": uploaded_file.name,
                    "type": (
                        "PDF" if uploaded_file.name.lower().endswith(".pdf") else "PPTX"
                    ),
                    "size": uploaded_file.size,
                    "chunks": len(chunks),
                }
            )
        else:
            st.warning(f"No readable text was found in {uploaded_file.name}.")

    if not all_chunks:
        progress.empty()
        st.error(
            "No readable text was found. If the PDF is scanned, "
            "it may need OCR first."
        )
        st.stop()

    try:
        progress.progress(80, text="Creating the searchable study index...")

        vector_db = Chroma.from_texts(
            texts=all_chunks,
            embedding=embedding_model,
            collection_name=f"study_{uuid.uuid4().hex[:12]}",
        )

        st.session_state.vector_db = vector_db
        st.session_state.retriever = build_retriever(vector_db)
        st.session_state.processed_files = signature
        st.session_state.documents = doc_info
        st.session_state.generated_quiz = None
        st.session_state.quiz_submitted = False
        st.session_state.quiz_answers = {}
        st.session_state.last_exam_answer = None

        log_activity(f"Processed {len(doc_info)} file(s) into {len(all_chunks)} chunks")

        progress.progress(100, text="Your study materials are ready!")
        progress.empty()

        st.success(
            f"Ready to study! Processed {len(uploaded_files)} file(s) "
            f"into {len(all_chunks)} text chunks."
        )

    except Exception:
        progress.empty()
        st.error(
            "Could not create the study index. "
            "Check your API connection and try again."
        )
        st.stop()


def render_upload_card():
    with st.container(key="upload_card"):
        render(
            """
            <div class="panel-head">
              <div class="panel-ico">☁️</div>
              <div>
                <div class="panel-title">Upload Your Study Materials</div>
                <div class="panel-sub">Drag and drop PDF or PPTX files here,
                or click to browse</div>
              </div>
            </div>
            """
        )

        uploaded_files = st.file_uploader(
            "Drop your files here or browse your computer",
            type=["pdf", "pptx"],
            accept_multiple_files=True,
            help="Supported formats: PDF and PowerPoint .pptx. "
            "Older .ppt files must be saved as .pptx first.",
            label_visibility="collapsed",
        )

        render(
            """
            <div class="supported">Supported: <span class="tag">PDF</span>
            <span class="tag">PPTX</span> <span style="opacity:.5">|</span>
            Turn your notes into interactive learning</div>
            """
        )

        if uploaded_files:
            process_files(uploaded_files)


def render_documents_card():
    docs = st.session_state.documents

    if docs:
        rows = ""
        for d in docs:
            kind = d["type"].lower()
            rows += f"""
            <div class="doc-row">
              <div class="doc-ico {kind}">{'PDF' if kind == 'pdf' else 'P'}</div>
              <div class="doc-meta">
                <b>{html_lib.escape(d['name'])}</b>
                <span>{d['type']} • {human_size(d['size'])} • {d['chunks']} chunks</span>
              </div>
              <span class="pill-ok">● Processed</span>
            </div>
            """
        body = rows
    else:
        body = """
        <div class="empty-docs">No documents yet.<br/>
        Upload a PDF or PPTX and it will appear here once processed.</div>
        """

    with st.container(key="docs_card"):
        render(
            f"""
            <div class="panel-row" style="margin-bottom:.4rem">
              <div class="panel-title">📑 Your Documents</div>
            </div>
            {body}
            """
        )


def has_materials() -> bool:
    return st.session_state.retriever is not None


def need_materials():
    st.info(
        "Upload your PDF or PPTX files first — open **Upload Documents** "
        "from the sidebar."
    )
    st.button(
        "Go to Upload Documents",
        on_click=go,
        args=("Upload Documents",),
    )


# --------------------------------------------------
# Hero
# --------------------------------------------------
def render_hero():
    n_docs = len(st.session_state.documents)

    if has_materials():
        ai_text = f"{n_docs} document(s) indexed and ready to study."
    else:
        ai_text = "Your study assistant is online and ready to help."

    render(
        """
        <div class="topbar">
          <div class="welcome"><div class="avatar">🎓</div>
          <div>Welcome<small>Student</small></div></div>
        </div>
        """
    )

    render(
        f"""
        <div class="hero">
          <div class="hero-grid">
            <div>
              <div class="hero-kicker">✦ &nbsp;YOUR PERSONAL AI LEARNING SPACE</div>
              <h1>Learn <span class="grad">smarter.</span><br/>
                  Remember <span class="grad2">longer.</span></h1>
              <p>Your personal AI learning companion. Transform your study
              materials into explanations, interactive practice, and
              exam-ready answers.</p>
              <div class="chips">
                <span class="chip">🧠 AI-powered learning</span>
                <span class="chip">📄 Your own study materials</span>
                <span class="chip">⏱️ Learn at your own pace</span>
              </div>
            </div>
            <div class="orb-wrap">
              <div class="orb"></div>
              <div class="orb-ring"></div>
              <div class="orb-ring two"></div>
              <span class="spark a"></span><span class="spark b"></span>
              <span class="spark c"></span>
            </div>
          </div>
          <div class="ai-card">
            <div class="ai-title">AI READY <span class="dot"></span></div>
            {ai_text}
          </div>
        </div>
        """
    )


def render_feature_cards():
    cols = st.columns(2, gap="medium")

    with cols[0]:
        with st.container(key="card_quiz"):
            render(
                """
                <div class="feat">
                  <div class="feat-ico">🎮</div>
                  <div>
                    <h3>Interactive Quiz</h3>
                    <p>Test your understanding with AI-generated questions,
                    instant feedback, and explanations.</p>
                  </div>
                </div>
                <div class="feat-tags">
                  <span class="feat-tag">🔢 5 or 10 questions</span>
                  <span class="feat-tag">⚡ Instant scoring</span>
                  <span class="feat-tag">💡 Detailed explanations</span>
                </div>
                """
            )
            st.button(
                "Start a quiz",
                key="home_quiz",
                use_container_width=True,
                on_click=go,
                args=("Interactive Quiz",),
            )

    with cols[1]:
        with st.container(key="card_exam"):
            render(
                """
                <div class="feat">
                  <div class="feat-ico violet">📝</div>
                  <div>
                    <h3>Exam Preparation</h3>
                    <p>Turn your study materials into structured 10-mark
                    answers or apply concepts to real scenarios.</p>
                  </div>
                </div>
                <div class="feat-tags">
                  <span class="feat-tag">📘 Straightforward answers</span>
                  <span class="feat-tag">🧩 Scenario-based answers</span>
                  <span class="feat-tag">⬇️ Download &amp; revise</span>
                </div>
                """
            )
            st.button(
                "Prepare an answer",
                key="home_exam",
                use_container_width=True,
                on_click=go,
                args=("10-Mark Preparation",),
            )


# --------------------------------------------------
# Pages
# --------------------------------------------------
def page_home():
    render_hero()

    left, right = st.columns([1.55, 1], gap="medium")
    with left:
        render_upload_card()
    with right:
        render_documents_card()

    st.write("")
    render_feature_cards()


def page_upload():
    st.header("📂 Upload Documents")
    st.caption(
        "Add your PDF or PowerPoint notes. They are indexed so every tool can use them."
    )

    left, right = st.columns([1.55, 1], gap="medium")
    with left:
        render_upload_card()
    with right:
        render_documents_card()


def page_ask():
    st.header("💬 Ask from Notes")
    st.write("Ask a question and get an answer based on your uploaded materials.")

    if not has_materials():
        need_materials()
        return

    with st.form("document_question_form"):
        question = st.text_input(
            "What would you like to understand?",
            placeholder="Example: Explain the main purpose of TCP",
        )
        ask_submitted = st.form_submit_button(
            "Find answer",
            use_container_width=True,
        )

    if ask_submitted:
        if not question.strip():
            st.warning("Type a question first.")
        else:
            with st.spinner("Searching your materials..."):
                try:
                    docs = st.session_state.retriever.invoke(question)
                    context = "\n\n".join(doc.page_content for doc in docs)
                    response = llm.invoke(build_prompt(context, question))

                    st.markdown("#### Answer")
                    st.write(response.content)

                    with st.expander("See the matching text from your files"):
                        for n, doc in enumerate(docs, start=1):
                            st.markdown(f"**Extract {n}**")
                            st.write(doc.page_content)

                    log_activity(f"Asked: {question.strip()[:70]}")

                except Exception:
                    st.error("Could not generate an answer. Please try again.")


def clear_quiz_widgets():
    for idx in range(10):
        key = f"quiz_answer_{idx}"
        if key in st.session_state:
            del st.session_state[key]


def page_quiz():
    st.header("🎮 Interactive Quiz")
    st.write(
        "Check your understanding with questions generated from your own "
        "study materials."
    )

    render(
        """
        <div class="hint-card">💡 <b>Study tip:</b> Try answering from
        memory before checking the explanations.</div>
        """
    )

    if not has_materials():
        need_materials()
        return

    with st.container(border=True):
        quiz_topic = st.text_input(
            "What topic do you want to practise?",
            placeholder="Example: TCP and UDP",
            key="quiz_topic",
        )

        quiz_count = st.select_slider(
            "Quiz length",
            options=[5, 10],
            value=5,
            key="quiz_count",
        )

        if st.button(
            "✨ Create my quiz",
            key="generate_quiz",
            use_container_width=True,
        ):
            if not quiz_topic.strip():
                st.warning("Enter a topic first.")
            else:
                with st.spinner("Building your quiz from the uploaded materials..."):
                    try:
                        docs = st.session_state.retriever.invoke(quiz_topic)
                        context = "\n\n".join(doc.page_content for doc in docs)

                        prompt = f"""
Create exactly {quiz_count} multiple-choice questions about: {quiz_topic}
Use ONLY the study context below. Do not use outside facts.
Return valid JSON only, without Markdown fences, using this structure:
{{
  "questions": [
    {{
      "question": "Question text",
      "options": {{
        "A": "Option A",
        "B": "Option B",
        "C": "Option C",
        "D": "Option D"
      }},
      "answer": "A",
      "explanation": "Short explanation supported by the context"
    }}
  ]
}}
Each question must have four non-empty options A-D, one correct answer A-D,
and a concise explanation. Avoid duplicate questions.

Study context:
{context}
"""

                        raw = llm.invoke(prompt).content

                        if not isinstance(raw, str):
                            raise ValueError("Unexpected response format")

                        cleaned = re.sub(
                            r"^```(?:json)?\s*",
                            "",
                            raw.strip(),
                            flags=re.IGNORECASE,
                        )
                        cleaned = re.sub(r"\s*```$", "", cleaned)

                        questions = json.loads(cleaned).get("questions")

                        if (
                            not isinstance(questions, list)
                            or len(questions) != quiz_count
                        ):
                            raise ValueError("Unexpected question count")

                        for item in questions:
                            options = item.get("options", {})
                            answer = str(item.get("answer", "")).strip().upper()

                            valid = (
                                bool(item.get("question"))
                                and set(options) == {"A", "B", "C", "D"}
                                and all(
                                    isinstance(options.get(k), str)
                                    and options[k].strip()
                                    for k in "ABCD"
                                )
                                and answer in {"A", "B", "C", "D"}
                                and bool(item.get("explanation"))
                            )

                            if not valid:
                                raise ValueError("Invalid question")

                            item["answer"] = answer

                        st.session_state.generated_quiz = questions
                        st.session_state.quiz_topic_used = quiz_topic.strip()
                        st.session_state.quiz_submitted = False
                        st.session_state.quiz_answers = {}
                        clear_quiz_widgets()
                        log_activity(
                            f"Generated a {quiz_count}-question quiz on "
                            f"“{quiz_topic.strip()}”"
                        )
                        st.rerun()

                    except Exception:
                        st.error(
                            "Quiz generation failed. Please try again "
                            "with a topic covered by your uploaded files."
                        )

    questions = st.session_state.generated_quiz

    if not questions:
        return

    st.markdown("### 📝 Your questions")
    st.caption(
        "Choose one option for each question. You can leave a "
        "question unanswered and submit when you're ready."
    )

    with st.form("quiz_answer_form"):
        answers = {}

        for index, item in enumerate(questions):
            st.markdown(f"**{index + 1} of {len(questions)} · {item['question']}**")

            opts = item["options"]

            answers[str(index)] = st.radio(
                "Select your answer",
                options=["A", "B", "C", "D"],
                index=None,
                format_func=lambda key, options=opts: f"{key}. {options[key]}",
                key=f"quiz_answer_{index}",
                horizontal=False,
            )

            if index < len(questions) - 1:
                st.divider()

        submitted = st.form_submit_button(
            "Submit answers",
            use_container_width=True,
        )

    if submitted:
        st.session_state.quiz_answers = answers
        st.session_state.quiz_submitted = True

        score = sum(
            1
            for i, item in enumerate(questions)
            if answers.get(str(i)) == item["answer"]
        )
        st.session_state.quiz_history.append(
            {
                "topic": st.session_state.quiz_topic_used,
                "score": score,
                "total": len(questions),
                "time": datetime.now().strftime("%d %b, %I:%M %p"),
            }
        )
        log_activity(
            f"Finished quiz on “{st.session_state.quiz_topic_used}”: "
            f"{score}/{len(questions)}"
        )

    if st.session_state.quiz_submitted:
        correct = 0

        answered = sum(
            value is not None for value in st.session_state.quiz_answers.values()
        )

        for index, item in enumerate(questions):
            chosen = st.session_state.quiz_answers.get(str(index))
            if chosen == item["answer"]:
                correct += 1

        st.markdown("### 📊 Your results")

        metric1, metric2, metric3 = st.columns(3)
        metric1.metric("Score", f"{correct}/{len(questions)}")
        metric2.metric("Correct", correct)
        metric3.metric("Answered", f"{answered}/{len(questions)}")

        st.progress(correct / len(questions))

        for index, item in enumerate(questions):
            chosen = st.session_state.quiz_answers.get(str(index))

            with st.expander(
                f"Question {index + 1} · "
                + ("Correct ✅" if chosen == item["answer"] else "Review this answer"),
                expanded=(chosen != item["answer"]),
            ):
                if chosen is None:
                    st.warning("You skipped this question.")
                elif chosen == item["answer"]:
                    st.success("Correct answer!")
                else:
                    st.error("Not quite — use the explanation to review.")

                st.write(
                    f"**Correct answer:** {item['answer']}. "
                    f"{item['options'][item['answer']]}"
                )
                st.write(f"**Why:** {item['explanation']}")

        if st.button(
            "🔄 Practise another topic",
            key="retry_quiz",
            use_container_width=True,
        ):
            st.session_state.generated_quiz = None
            st.session_state.quiz_submitted = False
            st.session_state.quiz_answers = {}
            clear_quiz_widgets()
            st.rerun()


def page_exam():
    st.header("📝 Exam Preparation · 10-Mark Answer")
    st.write(
        "Prepare a structured answer for a direct question or practise "
        "applying concepts to a scenario."
    )

    if not has_materials():
        need_materials()
        return

    answer_type = st.radio(
        "Choose your answer style",
        ["📘 Straightforward Answer", "🧩 Scenario-Based Answer"],
        horizontal=True,
        key="ten_mark_answer_type",
    )

    exam_topic = ""
    scenario = ""
    scenario_question = ""

    if answer_type == "📘 Straightforward Answer":
        render(
            """
            <div class="hint-card">📘 <b>Direct exam question</b><br/>
            Enter a topic or question and get an introduction, organized
            key points, and a conclusion.</div>
            """
        )

        exam_topic = st.text_area(
            "Your topic or question",
            placeholder="Example: Explain the TCP protocol and its main features.",
            key="ten_mark_topic",
            height=100,
        )

        search_query = exam_topic.strip()
        question_to_answer = exam_topic.strip()

    else:
        render(
            """
            <div class="hint-card">🧩 <b>Apply what you learned</b><br/>
            Describe the situation, then write the exact question you need
            to answer about it.</div>
            """
        )

        scenario = st.text_area(
            "1. Describe the scenario",
            placeholder=(
                "Example: A company needs to transfer important files "
                "reliably between two offices over a network."
            ),
            key="scenario_description",
            height=120,
        )

        scenario_question = st.text_area(
            "2. Enter the question about this scenario",
            placeholder="Example: Which transport protocol is appropriate, and why?",
            key="scenario_question",
            height=90,
        )

        search_query = f"{scenario.strip()}\n{scenario_question.strip()}".strip()
        question_to_answer = (
            f"Scenario:\n{scenario.strip()}\n\n"
            f"Question:\n{scenario_question.strip()}"
        )

    if st.button(
        "✨ Generate my 10-mark answer",
        key="generate_10_mark",
        use_container_width=True,
    ):
        if answer_type == "📘 Straightforward Answer" and not exam_topic.strip():
            st.warning("Enter your topic or question first.")

        elif answer_type == "🧩 Scenario-Based Answer" and not scenario.strip():
            st.warning("Describe the scenario first.")

        elif (
            answer_type == "🧩 Scenario-Based Answer" and not scenario_question.strip()
        ):
            st.warning("Enter the question about your scenario.")

        else:
            with st.spinner("Building your answer from your study materials..."):
                try:
                    docs = st.session_state.retriever.invoke(search_query)
                    context = "\n\n".join(doc.page_content for doc in docs)

                    prompt = f"""
You are a university exam preparation assistant.
Write a detailed answer suitable for a 10-mark exam.

Answer style: {answer_type}

Question or task:
{question_to_answer}

Requirements:
- Use only information supported by the study context.
- Do not invent facts or add unsupported outside knowledge.
- Begin with an appropriate introduction or definition.
- Use clear headings and numbered points where helpful.
- Explain relevant concepts in sufficient detail for a 10-mark answer.
- Add examples only when supported by the context.
- End with a concise conclusion.
- For scenario-based answers, explicitly analyse the given scenario
  and answer the exact question by applying relevant concepts.
- Do not just repeat general definitions.
- If the context lacks enough information, clearly say what is missing.

Study context:
{context}

Write the answer now.
"""

                    response = llm.invoke(prompt)
                    st.session_state.last_exam_answer = response.content
                    log_activity(
                        f"Generated a 10-mark answer ({answer_type[2:].strip()})"
                    )

                except Exception:
                    st.error("Could not generate the answer. Please try again.")

    if st.session_state.last_exam_answer:
        st.markdown("### 📄 Your answer")
        st.markdown(st.session_state.last_exam_answer)

        st.caption(
            "Tip: Read the answer once, then try explaining the main "
            "points in your own words."
        )

        st.download_button(
            "⬇️ Download answer as text",
            data=st.session_state.last_exam_answer,
            file_name="10_mark_exam_answer.txt",
            mime="text/plain",
            use_container_width=True,
        )


def page_progress():
    st.header("📊 My Study Progress")
    history = st.session_state.quiz_history

    if not history:
        st.info("Complete a quiz and your results will show up here.")
        return

    total_score = sum(h["score"] for h in history)
    total_q = sum(h["total"] for h in history)

    c1, c2, c3 = st.columns(3)
    c1.metric("Quizzes taken", len(history))
    c2.metric("Questions answered correctly", f"{total_score}/{total_q}")
    c3.metric("Average accuracy", f"{round(100 * total_score / total_q)}%")

    st.markdown("### Quiz history")
    for h in reversed(history):
        pct = h["score"] / h["total"]
        with st.container(border=True):
            st.markdown(
                f"**{h['topic'] or 'Untitled topic'}** · "
                f"{h['score']}/{h['total']} · {h['time']}"
            )
            st.progress(pct)


def page_activity():
    st.header("🕒 Recent Activity")
    activity = st.session_state.activity

    if not activity:
        st.info("Nothing yet. Your uploads, questions and quizzes will be listed here.")
        return

    for item in activity:
        st.markdown(f"**{item['time']}** — {item['text']}")


def page_settings():
    st.header("⚙️ Settings")

    with st.container(border=True):
        new_k = st.slider(
            "Text chunks retrieved per question",
            min_value=2,
            max_value=10,
            value=st.session_state.top_k,
            help="More chunks give the AI more context but can add noise.",
        )

        if new_k != st.session_state.top_k:
            st.session_state.top_k = new_k
            if st.session_state.vector_db is not None:
                st.session_state.retriever = build_retriever(st.session_state.vector_db)
            st.toast("Retrieval setting updated")

    with st.container(border=True):
        st.markdown("**Reset**")
        if st.button("🗑️ Clear all materials and progress", use_container_width=True):
            for key in (
                "retriever",
                "vector_db",
                "processed_files",
                "generated_quiz",
                "quiz_answers",
                "last_exam_answer",
            ):
                st.session_state[key] = defaults[key]
            st.session_state.documents = []
            st.session_state.quiz_history = []
            st.session_state.activity = []
            st.session_state.quiz_submitted = False
            clear_quiz_widgets()
            st.success("Everything has been cleared.")


# --------------------------------------------------
# Router
# --------------------------------------------------
PAGES = {
    "Home": page_home,
    "Upload Documents": page_upload,
    "Ask from Notes": page_ask,
    "Interactive Quiz": page_quiz,
    "10-Mark Preparation": page_exam,
    "My Study Progress": page_progress,
    "Recent Activity": page_activity,
    "Settings": page_settings,
}

PAGES.get(st.session_state.page, page_home)()

st.divider()
st.caption(
    "Study Hub · Use your uploaded notes as the source, practise actively, "
    "and review the explanations after each attempt."
)
